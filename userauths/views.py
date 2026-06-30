import random
import string
from typing import Any
from urllib import request
from django.shortcuts import get_object_or_404, render, redirect
from django.urls import reverse_lazy
from .utils import send_account_created_email, send_password_reset_otp_email, send_activation_resend_email, get_client_ip
from django.http import HttpResponse, JsonResponse
from django.template.loader import render_to_string
from userauths.forms import EditUserProfileForm, PasswordChangingForm, CreateUserProfileForm, MyProfileForm
from django.contrib.auth import login, logout, authenticate
from django.contrib import messages
from django.conf import settings 
from userauths.models import CustomUser
from datetime import datetime
from django.views.generic import ListView, DetailView,CreateView, DeleteView, UpdateView, TemplateView
from django.views import View, generic
from django.contrib.auth.forms import UserChangeForm, PasswordChangeForm
from django.contrib.auth.views import PasswordChangeView
from django.contrib.auth.mixins import LoginRequiredMixin
#User = settings.AUTH_USER_MODEL
from django.contrib.auth.decorators import login_required
from django.views.decorators.http import require_POST
from .models import *
from userauths.forms import *
from .models import CustomUser
from PB_Entreprise.models import UserProfile
from userauths.profile_helpers import (
    all_permission_groups,
    build_profile_page_context,
    ensure_admin_profile,
    get_user_profile,
    save_user_profile,
    apply_profile_edit_form,
    apply_my_profile_form,
    USER_TYPE_EMAIL_SUBJECTS,
)
# Create your views here.
# from .utils import generate_greeting, generate_goodbye
from django.contrib.auth.hashers import make_password
from django.contrib.auth import get_user_model
from django.db import transaction
from django.utils import timezone
from .forms import ChangePasswordForm

from userauths.forms import CustomPermissionForm, TypeCustomPermissionForm
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
import pandas as pd

CustomUser = get_user_model()

def _accounts_queryset():
    return (
        CustomUser.objects.filter(is_superuser=False)
        .select_related('profile')
        .prefetch_related('profile__gerant_voiture', 'custom_permissions')
        .order_by('-created_at')
    )


def _form_errors_messages(request, form, prefix=''):
    for field, errors in form.errors.items():
        for error in errors:
            label = prefix or field
            messages.error(request, f"Erreur dans {label}: {error}")

@login_required(login_url='/login/')
def register(request):
    """Page unique : création et gestion de tous les types de comptes."""
    actor = request.user
    form = UserProfileForm()

    if request.method == 'POST':
        form = UserProfileForm(request.POST)
        user_type = request.POST.get('user_type', '')
        if user_type in ('2', '3', '4') and actor.user_type != '1':
            messages.error(request, "Seuls les administrateurs peuvent créer ce type de compte.")
            return redirect('register')

        if user_type in ('2', '3', '4'):
            ensure_admin_profile(actor)

        if form.is_valid():
            try:
                with transaction.atomic():
                    new_user = CustomUser(
                        username=form.cleaned_data['username'],
                        email=form.cleaned_data['email'],
                        gender=form.cleaned_data['gender'],
                    )
                    password = generate_random_password()
                    ut = form.cleaned_data['user_type']
                    _create_user_with_password(new_user, password, ut)
                    create_by = actor if ut != '1' else None
                    save_user_profile(new_user, form, create_by=create_by)
                    new_user.custom_permissions.set(form.cleaned_data.get('permissions', []))

                    has_send = send_account_created_email(
                        user=new_user,
                        password=password,
                        subject=USER_TYPE_EMAIL_SUBJECTS.get(ut, 'Création de compte — P&BENTREPRISE'),
                    )
                if has_send:
                    messages.success(request, 'Compte créé avec succès. Un email a été envoyé.')
                else:
                    messages.warning(request, 'Compte créé avec succès, mais l\'email n\'a pas pu être envoyé.')
                return redirect('register')
            except Exception as e:
                messages.error(request, f"Erreur: {str(e)}")
        else:
            _form_errors_messages(request, form)

    selected_permission_ids = set()
    raw_permissions = form['permissions'].value()
    if raw_permissions:
        selected_permission_ids = {int(v) for v in raw_permissions}

    return render(request, 'register.html', {
        'form': form,
        'accounts': _accounts_queryset(),
        'permission_groups': all_permission_groups(),
        'selected_permission_ids': selected_permission_ids,
    })


def list_users(request):
    return register(request)


@login_required
def edit_user_permissions(request, user_id):
    user = get_object_or_404(CustomUser, id=user_id)
    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'
    
    if request.method == 'POST':
        form = UserPermissionForm(request.POST)
        if form.is_valid():
            permissions = form.cleaned_data['permissions']
            user.custom_permissions.set(permissions)
            messages.success(request, 'Permissions mises à jour avec succès.')
            if is_ajax:
                return JsonResponse({
                    'success': True,
                    'message': 'Permissions mises à jour avec succès.'
                })
            return redirect('edit_user_permissions', user.id)  
    else:
        form = UserPermissionForm(initial={
            'permissions': user.custom_permissions.all()
        })

    raw_permissions = form['permissions'].value()
    if raw_permissions:
        selected_permission_ids = {int(v) for v in raw_permissions}
    else:
        selected_permission_ids = set(user.custom_permissions.values_list('pk', flat=True))

    perm_context = {
        'form': form,
        'user': user,
        'permission_groups': all_permission_groups(),
        'selected_permission_ids': selected_permission_ids,
    }

    # Si c'est une requête AJAX, retourner seulement le contenu de la modal
    if is_ajax:
        html = render_to_string('modif_user_perm_modal.html', perm_context, request=request)
        return JsonResponse({'html': html})

    return render(request, 'modif_user_perm.html', perm_context)

def generate_random_password(length=8):
    characters = string.ascii_letters + string.digits 
    return ''.join(random.choice(characters) for i in range(length))


@login_required(login_url='/login/')
def edit_account(request, user_id):
    """Édition d'un compte (tous rôles) via modal AJAX."""
    user_obj = get_object_or_404(CustomUser, id=user_id, is_superuser=False)
    try:
        profile = user_obj.profile
    except UserProfile.DoesNotExist:
        profile, _ = UserProfile.objects.get_or_create(user=user_obj)

    is_ajax = request.headers.get('X-Requested-With') == 'XMLHttpRequest'

    if request.method == 'POST':
        form = UserProfileEditForm(request.POST, instance=profile)
        if form.is_valid():
            try:
                apply_profile_edit_form(user_obj, profile, form)
                if is_ajax:
                    return JsonResponse({'success': True, 'message': 'Compte modifié avec succès.'})
                messages.success(request, 'Compte modifié avec succès.')
                return redirect('register')
            except Exception as e:
                if is_ajax:
                    return JsonResponse({'success': False, 'message': f"Erreur: {str(e)}"})
                messages.error(request, f"Erreur: {str(e)}")
        elif is_ajax:
            html = render_to_string('edit_account_modal.html', {
                'form': form, 'account_user': user_obj, 'profile': profile,
            }, request=request)
            return JsonResponse({'html': html})
    else:
        form = UserProfileEditForm(instance=profile)

    if is_ajax:
        html = render_to_string('edit_account_modal.html', {
            'form': form, 'account_user': user_obj, 'profile': profile,
        }, request=request)
        return JsonResponse({'html': html})

    return render(request, 'edit_account.html', {
        'form': form, 'account_user': user_obj, 'profile': profile,
    })


@login_required(login_url='/login/')
def delete_account(request, user_id):
    try:
        user_obj = get_object_or_404(CustomUser, id=user_id, is_superuser=False)
        if user_obj == request.user:
            messages.error(request, "Vous ne pouvez pas supprimer votre propre compte.")
            return redirect('register')
        label = user_obj.username
        try:
            user_obj.profile.delete()
        except UserProfile.DoesNotExist:
            pass
        user_obj.delete()
        messages.success(request, f"Le compte {label} a été supprimé avec succès.")
    except Exception as e:
        messages.error(request, f"Erreur lors de la suppression : {str(e)}")
    return redirect('register')


def password_success(request):
    return render(request,'userauths/success.html')

def pb_home(request):
    return render(request,'no_acces.html')
    # return render(request,'perfect/pb_home.html')
    
MAX_LOGIN_ATTEMPTS = 3  # nombre d'échecs autorisés avant blocage automatique


def _create_user_with_password(user, password, user_type):
    user.user_type = user_type
    user.email_verified = False
    user.is_active = False
    user.set_password(password)
    user.save()
    PasswordHistory.objects.create(user=user, password_hash=user.password)


def _record_login_history(request, user, *, successful, failure_reason=''):
    ip = get_client_ip(request) or '127.0.0.1'
    LoginHistory.objects.create(
        user=user,
        ip_address=ip,
        user_agent=(request.META.get('HTTP_USER_AGENT') or '')[:500],
        login_successful=successful,
        failure_reason=failure_reason,
    )


def verify_email_view(request, token):
    try:
        token_obj = EmailVerificationToken.objects.select_related('user').get(token=token, used=False)
    except EmailVerificationToken.DoesNotExist:
        messages.error(request, "Lien d'activation invalide ou déjà utilisé.")
        return redirect('login')

    if not token_obj.is_valid():
        messages.error(request, "Le lien d'activation a expiré. Contactez l'administrateur pour en recevoir un nouveau.")
        return redirect('login')

    user = token_obj.user
    now = timezone.now()
    user.email_verified = True
    user.email_verified_at = now
    user.is_active = True
    user.save(update_fields=['email_verified', 'email_verified_at', 'is_active'])

    token_obj.used = True
    token_obj.used_at = now
    token_obj.save(update_fields=['used', 'used_at'])
    EmailVerificationToken.objects.filter(user=user, used=False).exclude(id=token_obj.id).update(
        used=True, used_at=now,
    )

    messages.success(request, "Votre compte a été activé avec succès. Vous pouvez maintenant vous connecter.")
    return redirect('login')


class ResendActivationView(View):
    def get(self, request):
        return render(request, 'perfect/resend_activation.html')

    def post(self, request):
        email = (request.POST.get('email') or '').strip()
        try:
            user = CustomUser.objects.get(email=email)
        except CustomUser.DoesNotExist:
            messages.error(request, "Aucun compte trouvé avec cette adresse email.")
            return render(request, 'perfect/resend_activation.html')

        if user.email_verified:
            messages.info(request, "Ce compte est déjà activé. Vous pouvez vous connecter.")
            return redirect('login')

        if send_activation_resend_email(user):
            messages.success(
                request,
                f"Un nouveau lien d'activation a été envoyé à {user.email}. Vérifiez votre boîte de réception.",
            )
            return redirect('login')

        messages.error(request, "L'email d'activation n'a pas pu être envoyé. Réessayez plus tard.")
        return render(request, 'perfect/resend_activation.html')


def loginview(request):
    if request.user.is_authenticated:
        messages.warning(request, "hey you are already logged In")
        return redirect("home")

    if request.method == "POST":
        email = (request.POST.get("email") or "").strip()
        password = request.POST.get("password") or ""

        # 1) On vérifie d'abord l'existence du compte par email.
        try:
            existing_user = CustomUser.objects.get(email=email)
        except CustomUser.DoesNotExist:
            existing_user = None

        # 2) Compte non activé par email.
        if existing_user is not None and not existing_user.email_verified:
            _record_login_history(
                request,
                existing_user,
                successful=False,
                failure_reason="Compte non activé par email",
            )
            messages.error(
                request,
                "Votre compte n'est pas encore activé. Veuillez cliquer sur le lien reçu par email.",
            )
            return render(request, "perfect/logins.html")

        # 3) Compte bloqué (administrateur ou tentatives infructueuses).
        if existing_user is not None and not existing_user.is_active:
            _record_login_history(
                request,
                existing_user,
                successful=False,
                failure_reason="Compte bloqué",
            )
            messages.error(
                request,
                "Votre compte est bloqué. Veuillez contacter l'administrateur de l'application.",
            )
            return render(request, "perfect/logins.html")

        # 4) Authentification.
        user = authenticate(request, email=email, password=password) if existing_user else None

        if user is not None:
            # Connexion réussie → on remet le compteur à zéro.
            if user.failed_login_attempts or user.date_blocage:
                CustomUser.objects.filter(pk=user.pk).update(
                    failed_login_attempts=0,
                    date_blocage=None,
                )
            user.last_activity = timezone.now()
            user.save(update_fields=['last_activity'])
            _record_login_history(request, user, successful=True)
            login(request, user)
            user_type = user.user_type
            if user_type == '1':
                messages.success(request, f"Bienvenue Administrateur {user.username}")
                return redirect('dash')
            elif user_type == '2':
                messages.success(request, f"Bienvenue Chef d'exploitation {user.username}")
                return redirect('dash')
            elif user_type == '3':
                messages.success(request, f"Bienvenue Comptable {user.username}")
                return redirect('dash')
            elif user_type == '4':
                messages.success(request, f"Bienvenue Gérant {user.username}")
                return redirect('dashgarage')
            else:
                return redirect('login')
        # 5) Échec d'authentification.
        if existing_user is not None:
            # Le compte existe → on incrémente le compteur d'échecs.
            existing_user.failed_login_attempts = (existing_user.failed_login_attempts or 0) + 1

            if existing_user.failed_login_attempts >= MAX_LOGIN_ATTEMPTS:
                existing_user.is_active = False
                existing_user.date_blocage = timezone.now()
                existing_user.save(update_fields=[
                    'failed_login_attempts', 'is_active', 'date_blocage',
                ])
                from userauths.notification_utils import notify_account_blocked
                notify_account_blocked(existing_user, reason='auto')
                _record_login_history(
                    request,
                    existing_user,
                    successful=False,
                    failure_reason="Blocage après 3 tentatives infructueuses",
                )
                messages.error(
                    request,
                    "Votre compte est bloqué après 3 tentatives infructueuses. "
                    "Veuillez contacter l'administrateur de l'application.",
                )
            else:
                existing_user.save(update_fields=['failed_login_attempts'])
                _record_login_history(
                    request,
                    existing_user,
                    successful=False,
                    failure_reason="Identifiants invalides",
                )
                restant = MAX_LOGIN_ATTEMPTS - existing_user.failed_login_attempts
                messages.error(
                    request,
                    f"Mot de passe ou email invalide. Il vous reste {restant} tentative(s) "
                    f"avant le blocage automatique du compte.",
                )
        else:
            # Aucun compte avec cet email → message générique (on n'expose pas l'info).
            messages.error(request, "Détails de connexion invalides!!!")

    return render(request, "perfect/logins.html")

def logout_view(request):
    user=request.user
    logout(request)
    messages.success(request, f"Vous êtes deconnecté {user.username}")
    return redirect("login")

@require_POST
@login_required(login_url="login")
def toggle_active_user(request, pk):
    from userauths.notification_utils import notify_account_blocked, clear_account_blocked_notification

    user = get_object_or_404(CustomUser, id=pk)
    user.is_active = not user.is_active
    update_fields = ['is_active']
    if not user.is_active:
        user.date_blocage = timezone.now()
        update_fields.append('date_blocage')
    else:
        user.failed_login_attempts = 0
        user.date_blocage = None
        update_fields.extend(['failed_login_attempts', 'date_blocage'])
    user.save(update_fields=update_fields)
    if not user.is_active:
        notify_account_blocked(user, reason='admin')
    else:
        clear_account_blocked_notification(user)
    return JsonResponse({"success": True, "is_active": user.is_active})    

def interneView(request):
    return render(request,"userauths/interne.html")
# Vue pour afficher le formulaire de saisie de l'email
class ForgotPasswordView(View):
    def get(self, request):
        # return render(request, 'saisie_otp.html')
        return render(request, 'perfect/forgot_password.html')    

class RequestEmailView(View):
    def post(self, request):
        email = request.POST.get('email', '').strip()
        try:
            user = CustomUser.objects.get(email=email)

            PasswordResetOTP.objects.filter(
                user=user,
                used=False,
            ).update(used=True, used_at=timezone.now())

            otp_obj = PasswordResetOTP.objects.create(
                user=user,
                ip_address=get_client_ip(request),
            )

            if not send_password_reset_otp_email(user, otp_obj.otp):
                messages.error(request, "L'email de réinitialisation n'a pas pu être envoyé.")
                return render(request, "perfect/forgot_password.html")

            for key in ('password_reset_otp_id', 'password_reset_email', 'password_reset_otp_verified'):
                request.session.pop(key, None)
            request.session['password_reset_email'] = email
            request.session.modified = True

            messages.success(
                request,
                f"Un code de vérification a été envoyé à {user.email}. Il est valide pendant 5 minutes.",
            )
            return redirect("otp")
        except CustomUser.DoesNotExist:
            messages.error(request, "Utilisateur non trouvé.")
            return render(request, "perfect/forgot_password.html")


def _clear_password_reset_session(request):
    for key in ('password_reset_otp_id', 'password_reset_email', 'password_reset_otp_verified'):
        request.session.pop(key, None)
    request.session.modified = True


def _otp_page_context(email):
    """Contexte page OTP : email + horodatage d'expiration du code actif."""
    ctx = {'email': email}
    if not email:
        return ctx
    try:
        user = CustomUser.objects.get(email=email)
        otp_obj = (
            PasswordResetOTP.objects.filter(user=user, used=False)
            .order_by('-created_at')
            .first()
        )
        if otp_obj and otp_obj.expires_at:
            ctx['otp_expires_at'] = otp_obj.expires_at.isoformat()
            ctx['otp_expired'] = timezone.now() >= otp_obj.expires_at
    except CustomUser.DoesNotExist:
        pass
    return ctx


class VerifyOtpView(View):
    def get(self, request):
        if not request.session.get('password_reset_otp_verified'):
            messages.error(request, "Veuillez d'abord vérifier votre code OTP.")
            return redirect('otp')
        return render(request, 'perfect/reinitialise.html')

    def post(self, request):
        if not request.session.get('password_reset_otp_verified'):
            messages.error(request, "Veuillez d'abord vérifier votre code OTP.")
            return redirect('otp')

        new_password = request.POST.get('new_password')
        confirm_password = request.POST.get('confirm_password')
        email = request.session.get('password_reset_email')
        otp_id = request.session.get('password_reset_otp_id')

        if not new_password or not confirm_password:
            messages.error(request, "Tous les champs sont obligatoires.")
            return render(request, "perfect/reinitialise.html")

        if new_password != confirm_password:
            messages.error(request, "Les mots de passe ne correspondent pas.")
            return render(request, "perfect/reinitialise.html")

        if not email or not otp_id:
            messages.error(request, "Session expirée. Veuillez recommencer le processus.")
            _clear_password_reset_session(request)
            return redirect('mot_passe_oublie')

        try:
            user = CustomUser.objects.get(email=email)
            otp_obj = PasswordResetOTP.objects.get(id=otp_id, user=user, used=False)
        except (CustomUser.DoesNotExist, PasswordResetOTP.DoesNotExist):
            messages.error(request, "Code OTP invalide. Veuillez recommencer le processus.")
            _clear_password_reset_session(request)
            return redirect('mot_passe_oublie')

        if not otp_obj.is_valid():
            messages.error(request, "Le code OTP a expiré. Veuillez demander un nouveau code.")
            _clear_password_reset_session(request)
            return redirect('mot_passe_oublie')

        with transaction.atomic():
            user.password = make_password(new_password)
            user.save(update_fields=['password'])
            PasswordHistory.objects.create(user=user, password_hash=user.password)
            otp_obj.used = True
            otp_obj.used_at = timezone.now()
            otp_obj.save(update_fields=['used', 'used_at'])
            PasswordResetOTP.objects.filter(
                user=user,
                used=False,
            ).exclude(id=otp_obj.id).update(used=True, used_at=timezone.now())

        from userauths.notification_utils import notify_account_password_reset
        notify_account_password_reset(user)

        _clear_password_reset_session(request)
        messages.success(request, 'Mot de passe réinitialisé avec succès.')
        return redirect('login')


@login_required(login_url='login')
def notifications_list_api(request):
    from userauths.notification_utils import notifications_payload
    return JsonResponse(notifications_payload(request.user))


@login_required(login_url='login')
@require_POST
def notification_mark_read_api(request, pk):
    from userauths.notification_utils import mark_notification_read, notifications_payload
    mark_notification_read(request.user, pk)
    payload = notifications_payload(request.user)
    return JsonResponse({'success': True, 'unread_count': payload['unread_count']})


@login_required(login_url='login')
@require_POST
def notification_mark_all_read_api(request):
    from userauths.notification_utils import mark_all_notifications_read, notifications_payload
    mark_all_notifications_read(request.user)
    payload = notifications_payload(request.user)
    return JsonResponse({'success': True, 'unread_count': payload['unread_count']})
class OptValid(View):
    def get(self, request):
        email = request.session.get('password_reset_email')
        if not email:
            messages.error(request, "Veuillez d'abord saisir votre adresse email.")
            return redirect('mot_passe_oublie')
        return render(request, 'perfect/otp.html', _otp_page_context(email))

    def post(self, request):
        email = request.session.get('password_reset_email')
        if not email:
            messages.error(request, "Session expirée. Veuillez recommencer le processus.")
            return redirect('mot_passe_oublie')

        otp = request.POST.get('otp', '').strip()
        if len(otp) != 6 or not otp.isdigit():
            messages.error(request, "Le code OTP doit contenir 6 chiffres.")
            return render(request, 'perfect/otp.html', _otp_page_context(email))

        try:
            user = CustomUser.objects.get(email=email)
            otp_obj = PasswordResetOTP.objects.filter(
                user=user,
                otp=otp,
                used=False,
            ).order_by('-created_at').first()

            if not otp_obj:
                messages.error(request, "Code OTP invalide. Veuillez vérifier et réessayer.")
                return render(request, 'perfect/otp.html', _otp_page_context(email))

            if not otp_obj.is_valid():
                otp_obj.increment_attempts()
                if otp_obj.attempts >= 5:
                    messages.error(request, "Trop de tentatives échouées. Veuillez demander un nouveau code.")
                    return redirect('mot_passe_oublie')
                if timezone.now() >= otp_obj.expires_at:
                    messages.error(request, "Le code OTP a expiré. Veuillez demander un nouveau code.")
                else:
                    messages.error(request, "Code OTP invalide ou expiré.")
                return render(request, 'perfect/otp.html', _otp_page_context(email))

            request.session['password_reset_otp_id'] = otp_obj.id
            request.session['password_reset_otp_verified'] = True
            request.session.modified = True
            messages.success(request, "Code OTP vérifié avec succès. Définissez votre nouveau mot de passe.")
            return redirect("verify_otp")

        except CustomUser.DoesNotExist:
            messages.error(request, "Utilisateur non trouvé.")
            _clear_password_reset_session(request)
            return redirect('mot_passe_oublie')
                     
class PasswordChangeView(PasswordChangeView):
    form_class = PasswordChangingForm
    template_name = 'profil.html'
    success_message = "Mot de passe réinitialisé avec succès👍✓✓"
    error_message = "Erreur de saisie ✘✘"
    profile_success_message = "Profil mis à jour avec succès."
    success_url = reverse_lazy('change_password')

    def _get_profile(self, user):
        profile = get_user_profile(user)
        if profile is None:
            profile, _ = UserProfile.objects.get_or_create(user=user)
        return profile

    def _build_context(self, user, form=None, profile_form=None):
        if form is None:
            form = self.get_form()
        if profile_form is None:
            profile_form = MyProfileForm(user=user, profile=self._get_profile(user))
        return {
            'form': form,
            'profile_form': profile_form,
            **build_profile_page_context(user),
        }

    def form_valid(self, form):
        reponse = super().form_valid(form)
        messages.success(self.request, self.success_message)
        return reponse

    def form_invalid(self, form):
        reponse = super().form_invalid(form)
        messages.error(self.request, self.error_message)
        return reponse

    def get(self, request, *args, **kwargs):
        user = get_object_or_404(CustomUser, id=request.user.id)
        return render(request, self.template_name, self._build_context(user))

    def post(self, request, *args, **kwargs):
        user = get_object_or_404(CustomUser, id=request.user.id)
        profile = self._get_profile(user)

        if 'save_profile' in request.POST:
            profile_form = MyProfileForm(request.POST, request.FILES, user=user, profile=profile)
            if profile_form.is_valid():
                apply_my_profile_form(user, profile, profile_form)
                messages.success(request, "Profil mis à jour avec succès.")
                return redirect('change_password')
            context = self._build_context(user, profile_form=profile_form)
            return render(request, self.template_name, context)

        form = self.get_form()
        if form.is_valid():
            return self.form_valid(form)
        context = self._build_context(user, form=form)
        return render(request, self.template_name, context)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = get_object_or_404(CustomUser, id=self.request.user.id)
        context.update(self._build_context(user))
        return context

class PasswordChangeDoneView(View):
    def get(self, request):
         return render(request, 'password_change_done.html')

# ==================== GESTION DES PERMISSIONS ====================
class PermissionListView(LoginRequiredMixin, ListView):
    model = CustomPermission
    template_name = 'perfect/permission.html'
    context_object_name = 'permissions'
    paginate_by = 20
    
    def get_queryset(self):
        queryset = CustomPermission.objects.select_related('categorie').all()
        search = self.request.GET.get('search', '')
        categorie_filter = self.request.GET.get('categorie', '')
        
        if search:
            queryset = queryset.filter(name__icontains=search)
        if categorie_filter:
            queryset = queryset.filter(categorie__id=categorie_filter)
        
        return queryset.order_by('id','categorie__categorie', 'name')
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['categories'] = TypeCustomPermission.objects.all()
        context['form'] = CustomPermissionForm()
        permissions = list(context['permissions'])
        for perm in permissions:
            edit_form = CustomPermissionForm(instance=perm)
            edit_form.fields['name'].widget.attrs['id'] = f'id_name_{perm.pk}'
            edit_form.fields['categorie'].widget.attrs['id'] = f'id_categorie_{perm.pk}'
            edit_form.fields['url'].widget.attrs['id'] = f'id_url_{perm.pk}'
            perm.edit_form = edit_form
        context['permissions'] = permissions
        return context

class PermissionCreateView(LoginRequiredMixin, CreateView):
    model = CustomPermission
    form_class = CustomPermissionForm
    template_name = 'perfect/permission.html'
    success_url = reverse_lazy('list_permissions')
    
    def form_valid(self, form):
        messages.success(self.request, 'Permission créée avec succès ✓✓')
        return super().form_valid(form)
    
    def form_invalid(self, form):
        messages.error(self.request, 'Erreur lors de la création de la permission ✘✘')
        return super().form_invalid(form)

class PermissionUpdateView(LoginRequiredMixin, UpdateView):
    model = CustomPermission
    form_class = CustomPermissionForm
    template_name = 'perfect/partials/permission_form.html'
    success_url = reverse_lazy('list_permissions')
    success_message = 'Permission modifiée avec succès ✓✓'
    
    def form_valid(self, form):
        messages.success(self.request, self.success_message)
        return super().form_valid(form)
    
    def form_invalid(self, form):
        messages.error(self.request, 'Erreur lors de la modification ✘✘')
        return super().form_invalid(form)

@login_required
def delete_permission(request, pk):
    permission = get_object_or_404(CustomPermission, pk=pk)
    permission.delete()
    messages.success(request, 'Permission supprimée avec succès ✓✓')
    return redirect('list_permissions')

class ExportPermissionExcelView(LoginRequiredMixin, View):
    def get(self, request, *args, **kwargs):
        permissions = CustomPermission.objects.select_related('categorie').all().order_by('categorie__categorie', 'name')
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Permissions"
        
        # En-tête
        headers = ['ID', 'Nom', 'Catégorie', 'URL']
        ws.append(headers)
        
        # Style des en-têtes
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Données
        for perm in permissions:
            ws.append([
                perm.id,
                perm.name,
                perm.categorie.categorie,
                perm.url
            ])
        
        # Ajuster la largeur des colonnes
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 30
        
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = 'attachment; filename="Permissions.xlsx"'
        wb.save(response)
        return response

class ImportPermissionExcelView(LoginRequiredMixin, View):
    def post(self, request, *args, **kwargs):
        if 'excel_file' not in request.FILES:
            messages.error(request, 'Aucun fichier sélectionné ✘✘')
            return redirect('list_permissions')
        
        try:
            file = request.FILES['excel_file']
            wb = openpyxl.load_workbook(file)
            ws = wb.active
            created_count = 0
            updated_count = 0
            errors = []
            # Ignorer la première ligne (en-têtes)
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]:  # Ignorer les lignes vides
                    continue
                try:
                    perm_id = row[0]
                    name = row[1]
                    categorie_name = row[2]
                    url = row[3]
                    # Récupérer ou créer la catégorie
                    categorie, _ = TypeCustomPermission.objects.get_or_create(
                        categorie=categorie_name
                    )
                    # Créer ou mettre à jour la permission
                    permission, created = CustomPermission.objects.update_or_create(
                        id=perm_id,
                        defaults={
                            'name': name,
                            'categorie': categorie,
                            'url': url
                        }
                    )
                    if created:
                        created_count += 1
                    else:
                        updated_count += 1
                        
                except Exception as e:
                    errors.append(f"Ligne {row}: {str(e)}")
            
            if created_count > 0 or updated_count > 0:
                messages.success(
                    request, 
                    f'Import réussi: {created_count} créé(s), {updated_count} mis à jour ✓✓'
                )
            if errors:
                messages.warning(request, f'Erreurs: {len(errors)} ligne(s) en erreur')
                
        except Exception as e:
            messages.error(request, f'Erreur lors de l\'import: {str(e)} ✘✘')
        
        return redirect('list_permissions')

# ==================== GESTION DES CATÉGORIES ====================

class CategorieListView(LoginRequiredMixin, ListView):
    model = TypeCustomPermission
    template_name = 'perfect/categorie.html'
    context_object_name = 'categories'
    
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['form'] = TypeCustomPermissionForm()
        categories = list(context['categories'])
        for cat in categories:
            edit_form = TypeCustomPermissionForm(instance=cat)
            edit_form.fields['categorie'].widget.attrs['id'] = f'id_categorie_{cat.pk}'
            cat.edit_form = edit_form
        context['categories'] = categories
        return context

class CategorieCreateView(LoginRequiredMixin, View):
    """Vue combinée pour créer des catégories via formulaire ou import Excel"""
    login_url = 'login'
    success_url = reverse_lazy('list_categories')
    
    def post(self, request, *args, **kwargs):
        # Vérifier si c'est un import Excel
        if request.FILES.get('excel_file'):
            return self.handle_excel_import(request)
        else:
            # Sinon, traiter comme un formulaire normal
            return self.handle_form_submit(request)
    
    def handle_form_submit(self, request):
        """Gère la soumission du formulaire classique"""
        form = TypeCustomPermissionForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, 'Catégorie créée avec succès ✓✓')
            return redirect(self.success_url)
        else:
            messages.error(request, 'Erreur lors de la création de la catégorie ✘✘')
            return redirect('list_categories')
    
    def handle_excel_import(self, request):
        """Gère l'import depuis Excel"""
        excel_file = request.FILES['excel_file']
        try:
            df = pd.read_excel(excel_file)
            df = df.fillna('')
            
            # Vérifier les colonnes requises
            # Accepter soit 'categorie' soit une seule colonne sans en-tête
            if 'categorie' not in df.columns:
                # Si pas de colonne nommée 'categorie', prendre la première colonne
                if len(df.columns) == 1:
                    df.columns = ['categorie']
                else:
                    messages.error(request, "Format attendu : une seule colonne nommée 'categorie' ou une colonne unique")
                    return redirect('list_categories')
            
            created_count = 0
            skipped_count = 0
            errors = []
            
            for index, row in df.iterrows():
                categorie_name = str(row.get('categorie', '')).strip()
                if not categorie_name:
                    continue
                
                try:
                    # Vérifier si la catégorie existe déjà
                    categorie, created = TypeCustomPermission.objects.get_or_create(
                        categorie=categorie_name
                    )
                    if created:
                        created_count += 1
                    else:
                        skipped_count += 1
                except Exception as e:
                    errors.append(f"Ligne {index + 2}: {str(e)}")
            
            if created_count > 0:
                messages.success(request, f"✅ {created_count} catégorie(s) créée(s) avec succès")
            if skipped_count > 0:
                messages.info(request, f"ℹ️ {skipped_count} catégorie(s) déjà existante(s)")
            if errors:
                messages.warning(request, f"⚠️ {len(errors)} erreur(s) rencontrée(s)")
            
            return redirect('list_categories')
        
        except Exception as e:
            messages.error(request, f"Erreur lors de l'importation : {str(e)} ✘✘")
            return redirect('list_categories')

class CategorieUpdateView(LoginRequiredMixin, UpdateView):
    model = TypeCustomPermission
    form_class = TypeCustomPermissionForm
    template_name = 'perfect/partials/categorie_perm_form.html'
    success_url = reverse_lazy('list_categories')
    success_message = 'Catégorie modifiée avec succès ✓✓'

class ExportCategoriesExcelView(LoginRequiredMixin, View):
    """Vue pour exporter les catégories de permissions au format Excel (une seule colonne)"""
    login_url = 'login'
    
    def get(self, request, *args, **kwargs):
        categories = TypeCustomPermission.objects.all().order_by('categorie')
        
        # Création du fichier Excel
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Catégories de Permissions"
        
        # En-tête (une seule colonne)
        headers = ["categorie"]
        ws.append(headers)
        
        # Style pour l'en-tête
        header_fill = PatternFill(start_color="06497C", end_color="06497C", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=12)
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        cell = ws["A1"]
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        ws.column_dimensions["A"].width = 30
        
        # Données (une seule colonne avec les noms des catégories)
        for cat in categories:
            ws.append([cat.categorie])
        
        # Alignement des cellules de données
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=1):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Réponse HTTP
        filename = f"Categories-Permissions-{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename={filename}'
        wb.save(response)
        return response


@login_required
def delete_categorie(request, pk):
    categorie = get_object_or_404(TypeCustomPermission, pk=pk)
    # Vérifier si la catégorie est utilisée
    if CustomPermission.objects.filter(categorie=categorie).exists():
        messages.error(request, 'Impossible de supprimer: cette catégorie contient des permissions ✘✘')
    else:
        categorie.delete()
        messages.success(request, 'Catégorie supprimée avec succès ✓✓')
    return redirect('list_categories')